"""Import a reviewed material directory into the content-addressed archive.

The command is a dry run unless --apply is supplied. With --move-source, source
files are removed only after the database transaction and reports succeed.
"""

import argparse
import calendar
import csv
import hashlib
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import xlrd
from pypdf import PdfReader
from sqlalchemy import select

from app.config import Settings
from app.db import connect, now, uid
from app.models import (
    Document,
    DocumentMaterial,
    DocumentMaterialInvestor,
    DocumentMaterialProduct,
    Investor,
    InvestorProduct,
    Manager,
    Membership,
    Product,
    User,
)
from app.security import audit


PLACEHOLDER_PRODUCTS = {
    "PENDING-ZDVALUE1": {
        "name": "吉余斩道价值1号私募证券投资基金",
        "aliases": ["吉余斩道价值1号", "斩道价值1号"],
    }
}

CUSTOM_ALIASES = {
    "SA2889": ["SARD55", "吉余商指陆号", "商指陆号"],
    "SAAK02": ["吉余李字壹号", "李字壹号", "李字1号"],
    "SAQX82": ["吉余商指伍号", "商指伍号"],
    "SAVH33": ["吉余牡丹", "牡丹基金"],
    "SAWK26": ["吉余宸锋金炜幸福一号", "幸福一号", "幸福1号"],
    "SAYB79": ["吉余美提斯2号", "美提斯2号"],
    "SAYW68": ["吉余全球易一号", "全球易一号", "全球易1号", "全球易 合同"],
    "SB7648": ["吉余杰伦叁号", "杰伦叁号", "杰伦3号"],
    "SBBA95": ["吉余天乙2号", "吉余天乙二号", "天乙2号", "天乙二号"],
    "SBDK57": ["吉余击壤一号", "击壤一号", "击壤1号"],
    "SVH342": ["吉余天乙1号", "吉余天乙一号", "天乙1号", "天乙一号"],
    "SZB360": ["吉余杰伦壹号", "杰伦壹号", "杰伦1号"],
}

SETTLEMENT_ACCOUNT_PRODUCTS = {
    "001080012938": "PENDING-ZDVALUE1",
    "00151000902756": "SVH342",
    "1000902756": "SVH342",
    "M521015381": "SVH342",
    "东吴证券户天乙一号": "SVH342",
}


@dataclass
class ProductInfo:
    code: str
    name: str
    id: str | None = None
    aliases: set[str] = field(default_factory=set)


@dataclass
class PlanItem:
    source: Path
    relative_path: str
    sha256: str
    size: int
    category: str
    title: str
    sensitivity: str
    business_date: str | None
    period_start: str | None
    period_end: str | None
    investor_name: str | None
    investor_type: str | None
    quarter_investors: list[tuple[str, str]]
    product_codes: list[str]
    confidence: str
    evidence: str
    document_id: str | None = None
    moved: bool = False


def normalize(value: str) -> str:
    value = value.upper().replace("壹", "一").replace("贰", "二").replace("叁", "三")
    value = value.replace("一号", "1号").replace("二号", "2号").replace("三号", "3号")
    value = value.replace("上海吉余私募基金管理有限公司", "")
    return re.sub(r"[\s_—–\-（）(){}【】\[\]·,.，。]", "", value)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return " ".join((page.extract_text() or "") for page in reader.pages[:5])[:120000]
        if suffix == ".docx":
            parts: list[str] = []
            char_count = 0
            with zipfile.ZipFile(path) as archive:
                members = [
                    name
                    for name in archive.namelist()
                    if name == "word/document.xml"
                    or name.startswith("word/header")
                    or name.startswith("word/footer")
                ]
                for member in members:
                    root = ElementTree.fromstring(archive.read(member))
                    for node in root.iter():
                        if node.tag.endswith("}t") and node.text:
                            value = node.text.strip()
                            if value:
                                parts.append(value)
                                char_count += len(value)
                        if char_count >= 120000:
                            return " ".join(parts)[:120000]
            return " ".join(parts)[:120000]
        if suffix == ".xlsx":
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            char_count = 0
            for sheet in workbook.worksheets:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = [str(value) for value in row[:60] if value is not None]
                    parts.extend(values)
                    char_count += sum(map(len, values))
                    if row_number >= 2500 or char_count > 120000:
                        break
            return " ".join(parts)[:120000]
        if suffix == ".xls":
            workbook = xlrd.open_workbook(path, on_demand=True)
            parts = []
            char_count = 0
            for sheet in workbook.sheets():
                for row_number in range(min(sheet.nrows, 2500)):
                    values = [str(value) for value in sheet.row_values(row_number)[:60] if value != ""]
                    parts.extend(values)
                    char_count += sum(map(len, values))
                    if char_count > 120000:
                        break
            return " ".join(parts)[:120000]
        if suffix == ".txt":
            raw = path.read_bytes()[:240000]
            for encoding in ("utf-8", "gb18030", "utf-16"):
                try:
                    return raw.decode(encoding)
                except UnicodeDecodeError:
                    continue
            return raw.decode("utf-8", errors="ignore")
        if suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                return " ".join(member.filename for member in archive.infolist())[:120000]
    except Exception:
        return ""
    return ""


def extract_quarter_investors(path: Path) -> list[tuple[str, str]]:
    if path.suffix.lower() != ".xlsx" or "投资者" not in path.name:
        return []
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    for sheet in workbook.worksheets:
        if sheet.title in {"模板填报说明", "数据字典", "示例"}:
            continue
        rows = sheet.iter_rows(values_only=True)
        header = None
        name_index = type_index = None
        for row_number, row in enumerate(rows, 1):
            values = [str(value).strip() if value is not None else "" for value in row]
            if "投资者名称" in values:
                header = values
                name_index = values.index("投资者名称")
                type_index = values.index("投资者类型") if "投资者类型" in values else None
                break
            if row_number >= 20:
                break
        if header is None or name_index is None:
            continue
        result = []
        for row in rows:
            values = list(row)
            name = str(values[name_index]).strip() if name_index < len(values) and values[name_index] is not None else ""
            if not name:
                continue
            raw_type = str(values[type_index]).strip() if type_index is not None and type_index < len(values) and values[type_index] is not None else ""
            result.append((name, investor_type(raw_type, "")))
        return result
    return []


def investor_type(raw_type: str, top_level: str) -> str:
    if "管理人" in raw_type or "员工跟投" in raw_type:
        return "manager_co_investment"
    if "私募基金产品" in raw_type or top_level == "产品投资者":
        return "fund_product"
    if "计划" in raw_type:
        return "asset_management"
    if "自然人" in raw_type or top_level == "个人投资者":
        return "individual"
    if top_level == "机构投资者" or any(word in raw_type for word in ("法人", "机构", "公司")):
        return "institution"
    return "other"


def directory_investor(relative: Path) -> tuple[str | None, str | None]:
    parts = relative.parts
    if not parts or parts[0] not in {"个人投资者", "产品投资者", "机构投资者"}:
        return None, None
    if len(parts) < 3:
        return None, None
    top = parts[0]
    folder = parts[1]
    if top == "个人投资者" and folder.startswith("排排网财富") and len(parts) >= 4:
        folder = re.sub(r"-\d+$", "", parts[2])
    folder = folder.replace("——合格投资者", "").replace("合格投资者资料", "")
    folder = folder.strip(" -_")
    return (folder or None), investor_type("", top)


def category_for(relative: Path) -> tuple[str, str]:
    top = relative.parts[0]
    name = relative.name.lower()
    if top == "季报":
        return "periodic_report", "季度报告或季度报送材料"
    if top == "结算单库":
        return "custody_account", "账户结算单"
    if any(word in name for word in ("份额", "确认函", "信息变更", "回访")):
        return "operation_evidence", "投资者运营凭证"
    if any(word in name for word in ("认购", "申购", "赎回", "风险揭示", "风险匹配", "合同")):
        return "subscription_redemption", "合同及认申购材料"
    if "周报" in name:
        return "periodic_report", "定期报告"
    return "investor_qualification", "投资者准入与适当性材料"


def dates_for(path: Path, category: str) -> tuple[str | None, str | None, str | None]:
    name = path.name
    candidates = re.findall(r"(?<!\d)(20\d{2})[-年]?([01]\d)[-月]?([0-3]\d)(?!\d)", name)
    business_date = None
    for year, month, day in candidates:
        try:
            business_date = datetime(int(year), int(month), int(day)).date().isoformat()
            break
        except ValueError:
            continue
    month_match = re.search(r"(?<!\d)(20\d{2})[-年]?([01]\d)(?!\d)", name)
    if category == "custody_account" and month_match:
        year, month = map(int, month_match.groups())
        start = f"{year:04d}-{month:02d}-01"
        end = f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"
        return business_date or end, start, end
    if category == "periodic_report" and business_date:
        value = datetime.fromisoformat(business_date).date()
        quarter_start_month = ((value.month - 1) // 3) * 3 + 1
        return business_date, f"{value.year:04d}-{quarter_start_month:02d}-01", business_date
    return business_date, None, None


def build_products(db, manager_id: str) -> dict[str, ProductInfo]:
    result = {}
    for product in db.scalars(select(Product).where(Product.manager_id == manager_id)):
        aliases = {product.code, product.name, product.name.replace("私募证券投资基金", "")}
        aliases.update(CUSTOM_ALIASES.get(product.code, []))
        result[product.code] = ProductInfo(product.code, product.name, product.id, aliases)
    for code, entry in PLACEHOLDER_PRODUCTS.items():
        result.setdefault(code, ProductInfo(code, entry["name"], None, set(entry["aliases"])))
    return result


def match_product_codes(text: str, products: dict[str, ProductInfo]) -> list[str]:
    normalized = normalize(text)
    matches = []
    for code, product in products.items():
        aliases = sorted(product.aliases | {product.name, code}, key=len, reverse=True)
        if any(len(normalize(alias)) >= 4 and normalize(alias) in normalized for alias in aliases):
            matches.append(code)
    # The base 美提斯 name is contained in 美提斯2号; prefer the more specific product.
    if "SAYB79" in matches and "SZT138" in matches:
        matches.remove("SZT138")
    return sorted(set(matches))


def title_for(path: Path, label: str, investor_name: str | None) -> str:
    stem = re.sub(r"[_-]?[a-f0-9]{24,}$", "", path.stem, flags=re.I).strip(" _-")
    if re.fullmatch(r"[a-f0-9]{20,}", stem, flags=re.I) or stem.isdigit():
        media = "音视频" if path.suffix.lower() in {".mp4", ".aac"} else "影像"
        return f"{investor_name or '未识别主体'}{media}材料"
    return stem[:500] or label


def load_confirmed_relationships(source: Path, products: dict[str, ProductInfo]):
    relationships: dict[str, set[str]] = defaultdict(set)
    confirmed_names: set[str] = set()
    quarter_rows: dict[str, list[tuple[str, str]]] = {}
    for path in sorted((source / "季报").glob("*")):
        if not path.is_file():
            continue
        product_codes = match_product_codes(path.name, products)
        investors = extract_quarter_investors(path)
        quarter_rows[str(path.relative_to(source))] = investors
        if len(product_codes) == 1:
            for name, _ in investors:
                relationships[name].add(product_codes[0])
                confirmed_names.add(name)
    candidate_path = Path("runtime/analysis/资料关联候选_20260910.csv")
    if candidate_path.exists():
        with candidate_path.open(encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                code = row.get("关联产品代码", "")
                name = row.get("主体名称", "").strip()
                if name and code in products:
                    relationships[name].add(code)
    return relationships, confirmed_names, quarter_rows


def build_plan(source: Path, products: dict[str, ProductInfo]) -> list[PlanItem]:
    relationships, _, quarter_rows = load_confirmed_relationships(source, products)
    plan = []
    cached_files = []
    for path in sorted(item for item in source.rglob("*") if item.is_file()):
        relative = path.relative_to(source)
        text = extract_text(path)
        digest = sha256_file(path)
        size = path.stat().st_size
        investor_name, inv_type = directory_investor(relative)
        cached_files.append((path, relative, text, digest, size, investor_name, inv_type))
        if investor_name:
            direct_codes = match_product_codes(path.name, products) or match_product_codes(text, products)
            if len(direct_codes) == 1:
                relationships[investor_name].update(direct_codes)

    for path, relative, text, digest, size, investor_name, inv_type in cached_files:
        category, label = category_for(relative)
        quarter_investors = quarter_rows.get(str(relative), [])
        filename_codes = match_product_codes(path.name, products)
        content_codes = match_product_codes(text, products) if not filename_codes else []
        codes = filename_codes or content_codes
        evidence = "文件名命中产品" if filename_codes else "正文命中产品" if content_codes else ""
        if relative.parts[0] == "结算单库" and not codes:
            normalized_path = normalize(str(relative))
            for account, code in SETTLEMENT_ACCOUNT_PRODUCTS.items():
                if normalize(account) in normalized_path:
                    codes = [code]
                    evidence = "结算账户映射"
                    break
        if investor_name and not codes:
            codes = sorted(relationships.get(investor_name, set()))
            if codes:
                evidence = "投资者已知产品关系"
        business_date, period_start, period_end = dates_for(path, category)
        confidence = "high" if filename_codes or quarter_investors else "medium" if codes else "unlinked"
        plan.append(
            PlanItem(
                source=path,
                relative_path=str(relative),
                sha256=digest,
                size=size,
                category=category,
                title=title_for(path, label, investor_name),
                sensitivity="investor_sensitive" if relative.parts[0] in {"个人投资者", "产品投资者", "机构投资者"} or quarter_investors else "standard",
                business_date=business_date,
                period_start=period_start,
                period_end=period_end,
                investor_name=investor_name,
                investor_type=inv_type,
                quarter_investors=quarter_investors,
                product_codes=codes,
                confidence=confidence,
                evidence=evidence or "仅完成类别归纳，产品待核对",
            )
        )
    return plan


def admin_actor(db, manager_id: str):
    return next(
        (
            user
            for user, membership in db.execute(
                select(User, Membership)
                .join(Membership, Membership.user_id == User.id)
                .where(Membership.manager_id == manager_id)
            )
            if "admin" in (membership.roles or [])
        ),
        None,
    )


def ensure_archive_file(source: Path, target: Path, expected_sha: str):
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        if sha256_file(target) != expected_sha:
            raise RuntimeError(f"archive hash mismatch: {target}")
        return
    temporary = target.with_name(target.name + ".importing-" + uid())
    with source.open("rb") as incoming, temporary.open("xb") as outgoing:
        shutil.copyfileobj(incoming, outgoing, length=4 * 1024 * 1024)
        outgoing.flush()
        os.fsync(outgoing.fileno())
    if sha256_file(temporary) != expected_sha:
        temporary.unlink()
        raise RuntimeError(f"copy verification failed: {source}")
    os.replace(temporary, target)
    target.chmod(0o440)


def write_reports(plan: list[PlanItem], report_dir: Path, summary: dict):
    report_dir.mkdir(parents=True, exist_ok=True)
    csv_path = report_dir / "待整理全量导入清单_20260911.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["原路径", "SHA-256", "大小", "资料类别", "资料标题", "投资者", "产品代码", "匹配置信度", "匹配依据", "资料ID", "源文件已迁移"])
        for item in plan:
            writer.writerow([item.relative_path, item.sha256, item.size, item.category, item.title, item.investor_name or "", "、".join(item.product_codes), item.confidence, item.evidence, item.document_id or "", "是" if item.moved else "否"])
    category_counts = Counter(item.category for item in plan)
    product_counts = Counter(code for item in plan for code in item.product_codes)
    unresolved = [item for item in plan if not item.product_codes]
    markdown_path = report_dir / "待整理全量导入报告_20260911.md"
    lines = [
        "# 待整理全量导入报告",
        "",
        f"- 文件总数：{len(plan)}",
        f"- 总大小：{sum(item.size for item in plan):,} 字节",
        f"- 已关联产品：{sum(1 for item in plan if item.product_codes)}",
        f"- 产品待核对：{len(unresolved)}",
        f"- 新建或复用投资者：{summary.get('investors', 0)}",
        f"- 建立投资者—产品关系：{summary.get('investor_product_links', 0)}",
        f"- 新增资料记录：{summary.get('documents', 0)}",
        f"- 跳过已导入原路径：{summary.get('skipped', 0)}",
        f"- 已迁移源文件：{sum(1 for item in plan if item.moved)}",
        "",
        "## 分类统计",
        "",
        "| 类别 | 文件数 |",
        "|---|---:|",
    ]
    lines.extend(f"| {category} | {count} |" for category, count in sorted(category_counts.items()))
    lines += ["", "## 产品关联统计", "", "| 产品代码 | 文件数 |", "|---|---:|"]
    lines.extend(f"| {code} | {count} |" for code, count in sorted(product_counts.items()))
    lines += ["", "## 产品待核对", "", "以下文件已完成资料类别归纳，但没有足够证据建立产品关系：", ""]
    lines.extend(f"- `{item.relative_path}`" for item in unresolved)
    lines += ["", "完整逐文件结果见 `待整理全量导入清单_20260911.csv`。"]
    markdown_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return csv_path, markdown_path


def apply_plan(plan: list[PlanItem], settings: Settings, manager_name: str | None, move_source: bool, report_dir: Path):
    engine, factory = connect(settings.database_url)
    summary = Counter()
    with factory() as db:
        managers = list(db.scalars(select(Manager).order_by(Manager.name)))
        manager = next((item for item in managers if item.name == manager_name), None) if manager_name else managers[0] if len(managers) == 1 else None
        if not manager:
            raise SystemExit("无法唯一确定管理人，请使用 --manager 指定")
        actor = admin_actor(db, manager.id)
        if not actor:
            raise SystemExit("当前管理人没有管理员账号，无法写入审计记录")
        products = {product.code: product for product in db.scalars(select(Product).where(Product.manager_id == manager.id))}
        for code, entry in PLACEHOLDER_PRODUCTS.items():
            if code not in products and any(code in item.product_codes for item in plan):
                product = Product(
                    manager_id=manager.id,
                    code=code,
                    name=entry["name"],
                    currency="CNY",
                    strategy="",
                    expected=False,
                    frequency="off",
                    lifecycle_status="archived",
                    lifecycle_reason="仅从历史结算单识别，备案编码待核对",
                    lifecycle_updated_at=now(),
                    lifecycle_updated_by=actor.id,
                )
                db.add(product)
                db.flush()
                products[code] = product
                audit(db, actor, manager.id, "product.created_from_material_import", product.id, {"code": code, "name": product.name, "status": "pending_code_review"})
                summary["placeholder_products"] += 1
        existing_investors = {investor.display_name: investor for investor in db.scalars(select(Investor).where(Investor.manager_id == manager.id))}
        confirmed_names = {name for item in plan for name, _ in item.quarter_investors}
        planned_investors = {}
        for item in plan:
            if item.investor_name:
                planned_investors.setdefault(item.investor_name, item.investor_type or "other")
            for name, kind in item.quarter_investors:
                planned_investors.setdefault(name, kind)
        for name, kind in planned_investors.items():
            investor = existing_investors.get(name)
            if not investor:
                investor = Investor(
                    manager_id=manager.id,
                    investor_type=kind,
                    display_name=name,
                    status="confirmed" if name in confirmed_names else "pending",
                    source="directory_reference",
                    notes="由待整理目录导入；仅保存主体名称和类别，不保存证件号码",
                    created_by=actor.id,
                )
                db.add(investor)
                db.flush()
                existing_investors[name] = investor
                audit(db, actor, manager.id, "investor.created_from_directory_import", investor.id, {"display_name": name, "investor_type": kind})
                summary["investors_created"] += 1
        relationship_status: dict[tuple[str, str], str] = {}
        for item in plan:
            names = [(item.investor_name, item.investor_type)] if item.investor_name else []
            names += item.quarter_investors
            for name, kind in names:
                if not name:
                    continue
                for code in item.product_codes:
                    status = "confirmed" if (name, kind) in item.quarter_investors else "pending"
                    key = (name, code)
                    if status == "confirmed" or key not in relationship_status:
                        relationship_status[key] = status
        existing_links = {(link.investor_id, link.product_id) for link in db.scalars(select(InvestorProduct).where(InvestorProduct.manager_id == manager.id))}
        for (name, code), status in relationship_status.items():
            investor = existing_investors.get(name)
            product = products.get(code)
            if investor and product and (investor.id, product.id) not in existing_links:
                db.add(InvestorProduct(manager_id=manager.id, investor_id=investor.id, product_id=product.id, status=status, created_by=actor.id))
                existing_links.add((investor.id, product.id))
                summary["investor_product_links"] += 1
        existing_paths = {
            (document.metadata_json or {}).get("directory_import", {}).get("relative_path")
            for document in db.scalars(select(Document).where(Document.manager_id == manager.id, Document.source == "directory_import"))
        }
        for item in plan:
            if item.relative_path in existing_paths:
                summary["skipped"] += 1
                continue
            key = f"{manager.id}/{item.sha256[:2]}/{item.sha256}"
            target = settings.storage / key
            ensure_archive_file(item.source, target, item.sha256)
            linked_products = [products[code] for code in item.product_codes if code in products]
            document = Document(
                manager_id=manager.id,
                product_id=linked_products[0].id if len(linked_products) == 1 else None,
                filename=item.source.name[:255],
                sha256=item.sha256,
                storage_key=key,
                size=item.size,
                media_type=mimetypes.guess_type(item.source.name)[0] or "application/octet-stream",
                source="directory_import",
                uploader_id=actor.id,
                received_at=datetime.fromtimestamp(item.source.stat().st_mtime, timezone.utc).isoformat(),
                metadata_json={
                    "directory_import": {
                        "relative_path": item.relative_path,
                        "confidence": item.confidence,
                        "evidence": item.evidence,
                    }
                },
            )
            db.add(document)
            db.flush()
            item.document_id = document.id
            material = DocumentMaterial(
                document_id=document.id,
                manager_id=manager.id,
                category=item.category,
                title=item.title,
                business_date=item.business_date,
                period_start=item.period_start,
                period_end=item.period_end,
                notes=f"原路径：{item.relative_path}；归纳依据：{item.evidence}",
                sensitivity=item.sensitivity,
                status="organized",
                confirmed_by=actor.id,
                revision=1,
            )
            db.add(material)
            db.flush()
            for product in linked_products:
                db.add(DocumentMaterialProduct(manager_id=manager.id, document_id=document.id, product_id=product.id))
            linked_investor_names = {name for name, _ in item.quarter_investors}
            if item.investor_name:
                linked_investor_names.add(item.investor_name)
            for name in sorted(linked_investor_names):
                investor = existing_investors.get(name)
                if investor:
                    db.add(DocumentMaterialInvestor(manager_id=manager.id, document_id=document.id, investor_id=investor.id))
            audit(db, actor, manager.id, "document.imported_from_directory", document.id, {"relative_path": item.relative_path, "sha256": item.sha256, "category": item.category, "product_codes": item.product_codes, "investor_count": len(linked_investor_names)})
            summary["documents"] += 1
        db.commit()
        summary["investors"] = len(planned_investors)
    engine.dispose()
    csv_path, markdown_path = write_reports(plan, report_dir, summary)
    if move_source:
        for item in plan:
            if item.document_id and item.source.exists():
                item.source.unlink()
                item.moved = True
        for directory in sorted({item.source.parent for item in plan}, key=lambda path: len(path.parts), reverse=True):
            try:
                directory.rmdir()
            except OSError:
                pass
        write_reports(plan, report_dir, summary)
    return summary, csv_path, markdown_path


def backup_sqlite(database_url: str, output: Path):
    if not database_url.startswith("sqlite:///"):
        raise SystemExit("自动备份仅支持当前 SQLite 开发库")
    source = Path(database_url.removeprefix("sqlite:///"))
    output.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(source) as incoming, sqlite3.connect(output) as outgoing:
        incoming.backup(outgoing)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--manager")
    parser.add_argument("--report-dir", type=Path, default=Path("runtime/analysis"))
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--move-source", action="store_true")
    args = parser.parse_args()
    if args.move_source and not args.apply:
        parser.error("--move-source 必须与 --apply 一起使用")
    source = args.source.resolve()
    if not source.is_dir():
        parser.error("来源目录不存在")
    settings = Settings()
    engine, factory = connect(settings.database_url)
    with factory() as db:
        managers = list(db.scalars(select(Manager)))
        manager = next((item for item in managers if item.name == args.manager), None) if args.manager else managers[0] if len(managers) == 1 else None
        if not manager:
            raise SystemExit("无法唯一确定管理人，请使用 --manager 指定")
        products = build_products(db, manager.id)
    engine.dispose()
    plan = build_plan(source, products)
    preview = {
        "files": len(plan),
        "bytes": sum(item.size for item in plan),
        "categories": Counter(item.category for item in plan),
        "linked_files": sum(bool(item.product_codes) for item in plan),
        "unlinked_files": sum(not item.product_codes for item in plan),
        "investor_files": sum(bool(item.investor_name or item.quarter_investors) for item in plan),
        "duplicate_hashes": len(plan) - len({item.sha256 for item in plan}),
    }
    print(json.dumps(preview, ensure_ascii=False, indent=2, default=dict))
    if not args.apply:
        write_reports(plan, args.report_dir, {})
        print("当前为预览；使用 --apply 写入，增加 --move-source 在成功后移除源文件。")
        return
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = Path("runtime/backups") / f"development-before-material-import-{stamp}.db"
    backup_sqlite(settings.database_url, backup)
    print(f"database backup: {backup}")
    summary, csv_path, markdown_path = apply_plan(plan, settings, args.manager, args.move_source, args.report_dir)
    print(json.dumps(dict(summary), ensure_ascii=False, indent=2))
    print(f"manifest: {csv_path}")
    print(f"report: {markdown_path}")


if __name__ == "__main__":
    main()
