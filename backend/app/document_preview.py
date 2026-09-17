"""Safe, read-only previews for archived source documents."""

import csv
import io
import zipfile
from datetime import date, datetime
from html import escape
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


MAX_PREVIEW_ROWS = 200
MAX_PREVIEW_COLUMNS = 40
MAX_PREVIEW_SHEETS = 8
MAX_ARCHIVE_FILES = 2_000
MAX_ARCHIVE_BYTES = 80 * 1024 * 1024


def _page(title: str, body: str, note: str = "") -> str:
    safe_title = escape(title)
    note_html = f'<p class="note">{escape(note)}</p>' if note else ""
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>{safe_title}</title>
  <style>
    :root {{ color-scheme: light; font-family: Inter, "PingFang SC", "Microsoft YaHei", sans-serif; }}
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; padding: 22px; color: #33382f; background: #f8f7f2; }}
    header {{ position: sticky; top: 0; z-index: 2; margin: -22px -22px 18px; padding: 16px 22px; border-bottom: 1px solid #dddcd3; background: #fffdf9ee; backdrop-filter: blur(8px); }}
    h1 {{ margin: 0; font-size: 16px; font-weight: 650; overflow-wrap: anywhere; }}
    h2 {{ margin: 24px 0 10px; color: #5a6253; font-size: 14px; }}
    .note {{ margin: 6px 0 0; color: #7c8075; font-size: 12px; }}
    .sheet {{ max-width: 100%; margin-bottom: 24px; overflow: auto; border: 1px solid #dddcd3; border-radius: 9px; background: white; }}
    table {{ border-collapse: collapse; min-width: 100%; font-size: 12px; }}
    td {{ min-width: 90px; max-width: 320px; padding: 8px 10px; border-right: 1px solid #ebeae4; border-bottom: 1px solid #ebeae4; vertical-align: top; overflow-wrap: anywhere; white-space: pre-wrap; }}
    tr:first-child td {{ position: sticky; top: 0; color: #4f5949; background: #f0f2eb; font-weight: 650; }}
    .text {{ padding: 18px; border: 1px solid #dddcd3; border-radius: 9px; background: white; font: 13px/1.8 ui-monospace, SFMono-Regular, Consolas, monospace; white-space: pre-wrap; overflow-wrap: anywhere; }}
    .message {{ display: grid; min-height: 55vh; place-items: center; color: #74796c; text-align: center; }}
  </style>
</head>
<body><header><h1>{safe_title}</h1>{note_html}</header>{body}</body>
</html>"""


def _archive(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    members = archive.infolist()
    if len(members) > MAX_ARCHIVE_FILES or sum(item.file_size for item in members) > MAX_ARCHIVE_BYTES:
        archive.close()
        raise ValueError("文件解压大小或文件数量超出预览限制")
    return archive


def _value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _table(title: str, rows) -> str:
    body = []
    for row in rows:
        cells = list(row)[:MAX_PREVIEW_COLUMNS]
        body.append("<tr>" + "".join(f"<td>{escape(_value(cell))}</td>" for cell in cells) + "</tr>")
    return f'<h2>{escape(title)}</h2><div class="sheet"><table>{"".join(body)}</table></div>'


def _spreadsheet(filename: str, data: bytes) -> tuple[str, bool]:
    suffix = Path(filename).suffix.lower()
    sections = []
    truncated = False
    if suffix in {".xlsx", ".xlsm"}:
        with _archive(data):
            pass
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet_index, sheet in enumerate(book):
                if sheet_index >= MAX_PREVIEW_SHEETS:
                    truncated = True
                    break
                rows = []
                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index >= MAX_PREVIEW_ROWS:
                        truncated = True
                        break
                    if len(row) > MAX_PREVIEW_COLUMNS:
                        truncated = True
                    rows.append(row)
                sections.append(_table(sheet.title, rows))
        finally:
            book.close()
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for sheet_index, sheet in enumerate(book.sheets()):
                if sheet_index >= MAX_PREVIEW_SHEETS:
                    truncated = True
                    break
                rows = []
                for row_index in range(min(sheet.nrows, MAX_PREVIEW_ROWS)):
                    row = []
                    for cell in sheet.row(row_index)[:MAX_PREVIEW_COLUMNS]:
                        row.append(
                            xlrd.xldate_as_datetime(cell.value, book.datemode)
                            if cell.ctype == xlrd.XL_CELL_DATE
                            else cell.value
                        )
                    rows.append(row)
                truncated = truncated or sheet.nrows > MAX_PREVIEW_ROWS or sheet.ncols > MAX_PREVIEW_COLUMNS
                sections.append(_table(sheet.name, rows))
        finally:
            book.release_resources()
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("gb18030", errors="replace")
        rows = []
        for index, row in enumerate(csv.reader(io.StringIO(text))):
            if index >= MAX_PREVIEW_ROWS:
                truncated = True
                break
            if len(row) > MAX_PREVIEW_COLUMNS:
                truncated = True
            rows.append(row)
        sections.append(_table("CSV", rows))
    return "".join(sections), truncated


def _office_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    with _archive(data) as archive:
        if suffix == ".docx":
            names = ["word/document.xml"]
        else:
            names = sorted(
                name
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )[:MAX_PREVIEW_SHEETS]
        paragraphs = []
        for name in names:
            if name not in archive.namelist():
                continue
            root = ElementTree.fromstring(archive.read(name))
            text_nodes = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
            if text_nodes:
                paragraphs.append(" ".join(text_nodes))
        text = "\n\n".join(paragraphs)[:200_000]
    return f'<div class="text">{escape(text or "文件中没有可提取的文字内容。")}</div>'


def render_document_preview(path: Path, filename: str) -> str:
    """Return a self-contained escaped HTML preview for non-browser document formats."""
    data = path.read_bytes()
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".xlsx", ".xlsm", ".xls", ".csv"}:
            body, truncated = _spreadsheet(filename, data)
            note = "只读预览最多显示前 8 个工作表、每表 200 行和 40 列。" if truncated else "只读预览；原件内容未修改。"
            return _page(filename, body, note)
        if suffix in {".docx", ".pptx"}:
            return _page(filename, _office_text(filename, data), "只读文字预览；版式请下载原件核对。")
        if suffix in {".txt", ".log", ".md", ".json", ".xml"}:
            text = data[:500_000].decode("utf-8", errors="replace")
            return _page(filename, f'<div class="text">{escape(text)}</div>', "只读文本预览。")
    except (ValueError, OSError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        return _page(filename, f'<div class="message"><p>无法生成预览：{escape(str(exc))}</p></div>')
    return _page(filename, '<div class="message"><p>浏览器暂不支持此格式的站内预览，请下载原件查看。</p></div>')
