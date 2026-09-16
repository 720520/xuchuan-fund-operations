import hashlib
import io
import os
import re
import warnings
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import or_, select, update

from .db import now
from .models import (
    Document,
    DocumentMaterial,
    DocumentMaterialInvestor,
    DocumentMaterialProduct,
    EffectiveNav,
    ExceptionTask,
    Investor,
    InvestorPositionSnapshot,
    InvestorProduct,
    InvestorShareEvent,
    MailAction,
    MailItem,
    MailItemProduct,
    Membership,
    NavRecord,
    ParseJob,
    Product,
    ReceiptExpectation,
    ShareClass,
    ValidationRule,
)
from .parsing import business_date, numeric, parse
from .security import audit


INVESTOR_POSITION_HEADERS = {
    "产品代码",
    "产品名称",
    "投资者名称",
    "份额日期",
    "持有份额",
}


def _clean_cell(value):
    return re.sub(r"\s+", "", str(value or "").strip())


def _position_date(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = str(value or "").strip()
    match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if not match:
        match = re.fullmatch(r"(20\d{2})(\d{2})(\d{2})", text)
    if not match:
        raise ValueError("份额日期格式无法识别")
    return date(*(int(part) for part in match.groups())).isoformat()


def _position_units(value):
    text = str(value or "").replace(",", "").strip()
    amount = Decimal(text)
    if not amount.is_finite() or amount < 0:
        raise ValueError("持有份额必须是非负数字")
    return amount


def _position_rows(content):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        preview = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True))
        for offset, values in enumerate(preview, 1):
            headers = [_clean_cell(value) for value in values]
            if not INVESTOR_POSITION_HEADERS.issubset(set(headers)):
                continue
            indexes = {header: headers.index(header) for header in INVESTOR_POSITION_HEADERS}
            rows = []
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=offset + 1, values_only=True), offset + 1
            ):
                if not any(value not in (None, "") for value in row):
                    continue
                def cell(header):
                    index = indexes[header]
                    return row[index] if index < len(row) else None
                rows.append(
                    {
                        "row_number": row_number,
                        "product_code": _clean_cell(cell("产品代码")),
                        "product_name": str(cell("产品名称") or "").strip(),
                        "investor_name": str(cell("投资者名称") or "").strip(),
                        "as_of_date": _position_date(cell("份额日期")),
                        "units": _position_units(cell("持有份额")),
                    }
                )
            return rows
    return None


def _base_product_name(value):
    name = _clean_cell(value)
    return re.sub(r"(?:[A-Za-zＡ-Ｚａ-ｚ]类)(?:份额)?$", "", name)


def _share_name(value):
    match = re.search(r"([A-Za-zＡ-Ｚａ-ｚ])类(?:份额)?$", _clean_cell(value))
    if not match:
        return "总"
    letter = match.group(1)
    letter = chr(ord(letter) - 65248) if "Ａ" <= letter <= "Ｚ" or "ａ" <= letter <= "ｚ" else letter
    return f"{letter.upper()}类"


def _investor_type(name):
    if "基金" in name or "资管计划" in name or "资产管理计划" in name:
        return "fund_product"
    if any(word in name for word in ("公司", "合伙", "中心", "协会", "银行", "信托")):
        return "institution"
    return "individual"


def _import_actor_id(db, manager_id):
    memberships = list(
        db.scalars(
            select(Membership)
            .where(Membership.manager_id == manager_id)
            .order_by(Membership.id)
        )
    )
    if not memberships:
        return None
    memberships.sort(
        key=lambda membership: (
            "admin" not in (membership.roles or []),
            "operator" not in (membership.roles or []),
            membership.id,
        )
    )
    return memberships[0].user_id


def import_investor_position_documents(db, settings, limit=50):
    """Import stable custodian holding tables as investor position snapshots.

    Only workbooks carrying the complete position header are accepted. Other
    subscription/redemption attachments remain archived and are never interpreted
    as balances.
    """
    candidates = list(
        db.execute(
            select(Document, MailItem)
            .join(MailItem, MailItem.document_id == Document.parent_id)
            .where(
                Document.filename.ilike("%.xlsx"),
                or_(
                    MailItem.title.contains("投资人份额"),
                    MailItem.title.contains("投资者份额"),
                    Document.filename.contains("投资人份额"),
                    Document.filename.contains("投资者份额"),
                ),
            )
            .order_by(Document.received_at)
            .limit(limit)
        )
    )
    imported_documents = 0
    for document, mail_item in candidates:
        existing_material = db.get(DocumentMaterial, document.id)
        if (
            existing_material
            and existing_material.organization_source == "automatic"
            and existing_material.material_type == "position_statement"
        ):
            continue
        path = settings.storage / document.storage_key
        if not path.is_file():
            continue
        content = path.read_bytes()
        if hashlib.sha256(content).hexdigest() != document.sha256:
            continue
        try:
            rows = _position_rows(content)
        except Exception as exc:
            rows = None
            if "份额" in mail_item.title or "份额" in document.filename:
                task(
                    db,
                    document.manager_id,
                    "investor_share",
                    f"investor_position_import:{document.id}",
                    {
                        "document_id": document.id,
                        "source_mail_item_id": mail_item.id,
                        "reason": "position_statement_unreadable",
                        "message": "托管份额附件无法读取，原件已保留，请核对附件格式。",
                        "error_type": type(exc).__name__,
                    },
                )
        if rows is None:
            continue
        actor_id = _import_actor_id(db, document.manager_id)
        if not actor_id:
            task(
                db,
                document.manager_id,
                "investor_share",
                f"investor_position_import:{document.id}",
                {
                    "document_id": document.id,
                    "source_mail_item_id": mail_item.id,
                    "reason": "manager_has_no_member",
                    "message": "牌照下没有可用于建立投资者份额记录的成员账号。",
                },
            )
            continue

        products = list(
            db.scalars(select(Product).where(Product.manager_id == document.manager_id))
        )
        products_by_code = {product.code: product for product in products}
        products_by_name = {}
        for product in products:
            products_by_name.setdefault(_base_product_name(product.name), []).append(product)
        investors = list(
            db.scalars(select(Investor).where(Investor.manager_id == document.manager_id))
        )
        investors_by_name = {}
        for investor in investors:
            investors_by_name.setdefault(_clean_cell(investor.display_name), []).append(investor)
        shares_by_product = {
            product.id: {
                share.name: share
                for share in db.scalars(
                    select(ShareClass).where(ShareClass.product_id == product.id)
                )
            }
            for product in products
        }
        imported = issues = duplicates = 0
        linked_products = set()
        linked_investors = set()
        dates = set()
        for record in rows:
            product = products_by_code.get(record["product_code"])
            if not product:
                possible = products_by_name.get(_base_product_name(record["product_name"]), [])
                product = possible[0] if len(possible) == 1 else None
            if not product:
                issues += 1
                task(
                    db,
                    document.manager_id,
                    "investor_share",
                    f"investor_position_product:{document.id}:{record['row_number']}",
                    {
                        "document_id": document.id,
                        "source_mail_item_id": mail_item.id,
                        "reason": "product_not_found",
                        "message": "托管份额表中的产品尚未建档，完成产品建档后可重新导入。",
                        "row_number": record["row_number"],
                        "product_code": record["product_code"],
                    },
                )
                continue
            share_label = _share_name(record["product_name"])
            share = shares_by_product.get(product.id, {}).get(share_label)
            if not share:
                share = ShareClass(
                    manager_id=document.manager_id,
                    product_id=product.id,
                    name=share_label,
                )
                db.add(share)
                db.flush()
                shares_by_product.setdefault(product.id, {})[share_label] = share
                audit(
                    db,
                    None,
                    document.manager_id,
                    "share.auto_created_from_position_statement",
                    share.id,
                    {"product_id": product.id, "source_document_id": document.id},
                )
            normalized_name = _clean_cell(record["investor_name"])
            matches = investors_by_name.get(normalized_name, [])
            if len(matches) > 1:
                issues += 1
                task(
                    db,
                    document.manager_id,
                    "investor_share",
                    f"investor_position_investor:{document.id}:{record['row_number']}",
                    {
                        "document_id": document.id,
                        "source_mail_item_id": mail_item.id,
                        "product_id": product.id,
                        "reason": "investor_name_ambiguous",
                        "message": "托管份额表中的投资者名称对应多条档案，请合并重复投资者后重新导入。",
                        "row_number": record["row_number"],
                    },
                    product.id,
                    share.id,
                    record["as_of_date"],
                )
                continue
            if matches:
                investor = matches[0]
            else:
                investor = Investor(
                    manager_id=document.manager_id,
                    investor_type=_investor_type(record["investor_name"]),
                    display_name=record["investor_name"],
                    status="pending",
                    source="material",
                    notes="由托管投资人份额表自动建立，基础资料待补充。",
                    created_by=actor_id,
                )
                db.add(investor)
                db.flush()
                investors_by_name[normalized_name] = [investor]
                audit(
                    db,
                    None,
                    document.manager_id,
                    "investor.auto_created_from_position_statement",
                    investor.id,
                    {"source_document_id": document.id},
                )
            link = db.scalar(
                select(InvestorProduct).where(
                    InvestorProduct.investor_id == investor.id,
                    InvestorProduct.product_id == product.id,
                )
            )
            if not link:
                db.add(
                    InvestorProduct(
                        manager_id=document.manager_id,
                        investor_id=investor.id,
                        product_id=product.id,
                        created_by=actor_id,
                    )
                )
            existing = list(
                db.scalars(
                    select(InvestorPositionSnapshot).where(
                        InvestorPositionSnapshot.investor_id == investor.id,
                        InvestorPositionSnapshot.product_id == product.id,
                        InvestorPositionSnapshot.share_id == share.id,
                        InvestorPositionSnapshot.as_of_date == record["as_of_date"],
                    )
                )
            )
            if any(snapshot.units == record["units"] for snapshot in existing):
                duplicates += 1
            elif existing:
                issues += 1
                task(
                    db,
                    document.manager_id,
                    "investor_share",
                    f"investor_position_conflict:{investor.id}:{product.id}:{share.id}:{record['as_of_date']}",
                    {
                        "document_id": document.id,
                        "source_mail_item_id": mail_item.id,
                        "investor_id": investor.id,
                        "product_id": product.id,
                        "share_id": share.id,
                        "reason": "same_day_position_conflict",
                        "message": "同一投资者、产品和份额日期出现不同托管份额，未覆盖现有快照。",
                        "incoming_units": str(record["units"]),
                    },
                    product.id,
                    share.id,
                    record["as_of_date"],
                )
            else:
                snapshot = InvestorPositionSnapshot(
                    manager_id=document.manager_id,
                    investor_id=investor.id,
                    product_id=product.id,
                    share_id=share.id,
                    as_of_date=record["as_of_date"],
                    units=record["units"],
                    source_mail_item_id=mail_item.id,
                    source_document_id=document.id,
                    notes="托管投资人份额表自动导入",
                    created_by=actor_id,
                )
                db.add(snapshot)
                imported += 1
            linked_products.add(product.id)
            linked_investors.add(investor.id)
            dates.add(record["as_of_date"])
        timestamp = now()
        if mail_item.category != "investor_redemption" or mail_item.handling_mode == "pending":
            mail_item.category = "investor_redemption"
            mail_item.handling_mode = "receipt"
            mail_item.classification_source = "rule"
            mail_item.confidence = max(mail_item.confidence, 98)
            mail_item.status = "received"
            mail_item.updated_at = timestamp
            mail_item.revision += 1
        material = db.get(DocumentMaterial, document.id)
        if not material:
            material = DocumentMaterial(
                document_id=document.id,
                manager_id=document.manager_id,
                revision=1,
            )
            db.add(material)
        material.category = "subscription_redemption"
        material.material_type = "position_statement"
        material.title = document.filename
        material.business_date = next(iter(dates)) if len(dates) == 1 else None
        material.notes = f"托管持仓份额表自动导入：新增 {imported} 条，重复 {duplicates} 条，异常 {issues} 条。"
        material.sensitivity = "investor_sensitive"
        material.status = "organized"
        material.organization_source = "automatic"
        material.confirmed_by = actor_id
        material.confirmed_at = timestamp
        material.created_at = material.created_at or timestamp
        material.updated_at = timestamp
        for product_id in linked_products:
            if not db.scalar(
                select(DocumentMaterialProduct).where(
                    DocumentMaterialProduct.document_id == document.id,
                    DocumentMaterialProduct.product_id == product_id,
                )
            ):
                db.add(
                    DocumentMaterialProduct(
                        manager_id=document.manager_id,
                        document_id=document.id,
                        product_id=product_id,
                        created_at=timestamp,
                    )
                )
            if not db.scalar(
                select(MailItemProduct).where(
                    MailItemProduct.mail_item_id == mail_item.id,
                    MailItemProduct.product_id == product_id,
                )
            ):
                db.add(
                    MailItemProduct(
                        manager_id=document.manager_id,
                        mail_item_id=mail_item.id,
                        product_id=product_id,
                    )
                )
        for investor_id in linked_investors:
            if not db.scalar(
                select(DocumentMaterialInvestor).where(
                    DocumentMaterialInvestor.document_id == document.id,
                    DocumentMaterialInvestor.investor_id == investor_id,
                )
            ):
                db.add(
                    DocumentMaterialInvestor(
                        manager_id=document.manager_id,
                        document_id=document.id,
                        investor_id=investor_id,
                        created_at=timestamp,
                    )
                )
        action = db.scalar(select(MailAction).where(MailAction.mail_item_id == mail_item.id))
        if action and imported and not issues:
            action.status = "completed"
            action.result = {
                "object_type": "investor_position_snapshot_batch",
                "object_id": document.id,
                "imported_rows": imported,
                "duplicate_rows": duplicates,
            }
            action.updated_at = timestamp
            action.revision += 1
        job = db.scalar(select(ParseJob).where(ParseJob.document_id == document.id))
        if job and not (job.result or {}).get("record_ids"):
            job.status = "skipped"
            job.result = {
                **(job.result or {}),
                "skip_reason": "投资者持仓份额表已导入份额台账，不进入净值解析",
            }
            job.updated_at = timestamp
        audit(
            db,
            None,
            document.manager_id,
            "investor_position_statement.imported",
            document.id,
            {
                "source_mail_item_id": mail_item.id,
                "imported_rows": imported,
                "duplicate_rows": duplicates,
                "issue_rows": issues,
            },
        )
        imported_documents += 1
    return imported_documents


def task(
    db,
    manager_id,
    kind,
    dedup_key,
    payload,
    product_id=None,
    share_id=None,
    valuation_date=None,
):
    current = db.scalar(
        select(ExceptionTask)
        .where(
            ExceptionTask.manager_id == manager_id, ExceptionTask.dedup_key == dedup_key
        )
        .with_for_update()
    )
    if current:
        if current.status == "resolved":
            # New evidence reopens the same issue, preserving prior resolution in audit.
            audit(
                db,
                None,
                manager_id,
                "exception.reopened",
                current.id,
                {"previous_resolution": current.resolution},
            )
            current.status, current.assignee_id, current.resolution = "open", None, None
        current.payload = {**current.payload, **payload}
        current.updated_at, current.revision = now(), current.revision + 1
        return current
    current = ExceptionTask(
        manager_id=manager_id,
        kind=kind,
        dedup_key=dedup_key,
        payload=payload,
        product_id=product_id,
        share_id=share_id,
        valuation_date=valuation_date,
    )
    db.add(current)
    db.flush()
    return current


def reconcile_pending_investor_share_events(db, limit=200):
    """Upgrade legacy pending share rows to the email-first workflow."""
    events = list(
        db.scalars(
            select(InvestorShareEvent)
            .where(InvestorShareEvent.status == "pending")
            .order_by(InvestorShareEvent.created_at)
            .limit(limit)
            .with_for_update(skip_locked=True)
        )
    )
    timestamp = now()
    for event in events:
        if event.evidence_stage != "confirmation":
            event.status = "archived"
            event.updated_at = timestamp
            event.revision += 1
            audit(
                db,
                None,
                event.manager_id,
                "investor_share_event.auto_archived",
                event.id,
                {"evidence_stage": event.evidence_stage},
            )
            continue
        complete_email_record = bool(
            event.source_mail_item_id
            and (event.units_delta is not None or event.balance_after is not None)
            and (event.confirmation_date or event.effective_date)
        )
        if complete_email_record:
            event.status = "confirmed"
            event.confirmed_at = timestamp
            event.updated_at = timestamp
            event.revision += 1
            audit(
                db,
                None,
                event.manager_id,
                "investor_share_event.auto_confirmed",
                event.id,
                {"source_mail_item_id": event.source_mail_item_id},
            )
            continue
        task(
            db,
            event.manager_id,
            "investor_share",
            f"investor_share:{event.id}",
            {
                "share_event_id": event.id,
                "investor_id": event.investor_id,
                "source_mail_item_id": event.source_mail_item_id,
                "document_id": event.source_document_id,
                "reason": "confirmation_incomplete"
                if event.source_mail_item_id
                else "source_mail_missing",
                "message": "正式确认信息缺少可自动入账的数据或邮件依据，请核对来源。",
                "missing_units": event.units_delta is None
                and event.balance_after is None,
                "missing_mail": event.source_mail_item_id is None,
            },
            event.product_id,
            event.share_id,
            event.effective_date or event.confirmation_date,
        )
    return len(events)


def archive(
    db,
    settings,
    manager_id,
    name,
    content,
    source,
    actor=None,
    product_id=None,
    parent_id=None,
    received_at=None,
    metadata=None,
):
    if len(content) > settings.max_upload:
        limit = settings.max_upload // (1024 * 1024)
        raise HTTPException(413, f"文件超出 {limit} MiB 限制")
    if not content:
        raise HTTPException(422, "文件为空")
    name = name.replace("\\", "/").rsplit("/", 1)[-1][:255] or "attachment.bin"
    sha = hashlib.sha256(content).hexdigest()
    key = f"{manager_id}/{sha[:2]}/{sha}"
    path = settings.storage / key
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with path.open("xb") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        path.chmod(0o440)
    except FileExistsError:
        if hashlib.sha256(path.read_bytes()).hexdigest() != sha:
            raise HTTPException(500, "归档完整性异常")
    doc = Document(
        manager_id=manager_id,
        product_id=product_id,
        filename=name,
        sha256=sha,
        storage_key=key,
        size=len(content),
        media_type="message/rfc822"
        if name.lower().endswith(".eml")
        else "application/octet-stream",
        source=source,
        uploader_id=actor.id if actor else None,
        parent_id=parent_id,
        received_at=received_at or now(),
        metadata_json=metadata or {},
    )
    db.add(doc)
    db.flush()
    audit(
        db,
        actor,
        manager_id,
        "document.archived",
        doc.id,
        {"filename": name, "sha256": sha, "source": source},
    )
    return doc


def change_product_lifecycle(
    db,
    settings,
    product,
    actor,
    status,
    effective_date,
    reason,
    material_name=None,
    material_content=None,
):
    allowed = {"active", "liquidating", "liquidated", "archived"}
    if status not in allowed:
        raise HTTPException(422, "产品生命周期状态无效")
    reason = reason.strip()
    if not reason or len(reason) > 2000:
        raise HTTPException(422, "状态变更原因必填且不能超过 2000 字")
    if status == product.lifecycle_status:
        raise HTTPException(422, "产品已经处于该状态")
    if status == "liquidated" and not effective_date:
        raise HTTPException(422, "标记已清算必须填写清算完成日期")
    if status == "liquidated" and material_content is None:
        raise HTTPException(422, "标记已清算必须上传清算报告或托管确认材料")
    if isinstance(effective_date, str):
        effective_date = date.fromisoformat(effective_date)
    today = datetime.now(ZoneInfo("Asia/Shanghai")).date()
    if effective_date and effective_date > today:
        raise HTTPException(422, "状态生效日期不能晚于今天")
    material_doc = None
    if material_content is not None:
        material_doc = archive(
            db,
            settings,
            product.manager_id,
            material_name or "清算材料.bin",
            material_content,
            "lifecycle_material",
            actor,
            product.id,
            metadata={"lifecycle_status": status, "reason": reason},
        )
    before = {
        "status": product.lifecycle_status,
        "date": product.lifecycle_date,
        "reason": product.lifecycle_reason,
        "expected": product.expected,
        "frequency": product.frequency,
    }
    product.lifecycle_status = status
    product.lifecycle_date = (effective_date or today).isoformat()
    product.lifecycle_reason = reason
    product.lifecycle_updated_at = now()
    product.lifecycle_updated_by = actor.id
    if status in {"liquidated", "archived"}:
        db.execute(update(ReceiptExpectation).where(ReceiptExpectation.product_id == product.id).values(cancelled=True))
        product.expected = False
        product.frequency = "off"
        for issue in db.scalars(
            select(ExceptionTask).where(
                ExceptionTask.product_id == product.id,
                ExceptionTask.kind == "missing",
                ExceptionTask.status != "resolved",
            )
        ):
            issue.status = "resolved"
            issue.updated_at = now()
            issue.revision += 1
            issue.resolution = {
                "via": "product_lifecycle_changed",
                "status": status,
                "reason": reason,
            }
    audit(
        db,
        actor,
        product.manager_id,
        "product.lifecycle_changed",
        product.id,
        {
            "before": before,
            "after": {
                "status": status,
                "date": product.lifecycle_date,
                "reason": reason,
                "expected": product.expected,
                "frequency": product.frequency,
            },
            "material_document_id": material_doc.id if material_doc else None,
        },
    )
    return material_doc


def validate_nav(values):
    errors = []
    if Decimal(values["unit_nav"]) <= 0:
        errors.append(
            {
                "rule": "unit_nav_positive",
                "message": "单位净值必须大于零",
                "overridable": False,
            }
        )
    if (
        values["valuation_date"]
        > datetime.now(ZoneInfo("Asia/Shanghai")).date().isoformat()
    ):
        errors.append(
            {
                "rule": "future_date",
                "message": "估值日期晚于当前日期",
                "overridable": False,
            }
        )
    for key in ["net_assets", "total_shares"]:
        if values.get(key) is not None and Decimal(values[key]) < 0:
            errors.append(
                {
                    "rule": key + "_nonnegative",
                    "message": "规模或份额不能为负数",
                    "overridable": False,
                }
            )
    return errors


def add_nav(
    db,
    manager_id,
    product,
    share,
    values,
    actor=None,
    document=None,
    row_key=None,
    manual=False,
):
    # Serialize candidates and effective pointer per share. PostgreSQL lock is the production guarantee.
    db.scalar(select(ShareClass).where(ShareClass.id == share.id).with_for_update())
    values = {**values, "valuation_date": business_date(values["valuation_date"])}
    for key in ["unit_nav", "accumulated_nav", "net_assets", "total_shares"]:
        values[key] = numeric(values.get(key), key)
    if not values["unit_nav"]:
        raise HTTPException(422, "缺少单位净值")
    if document and row_key:
        prior = db.scalar(
            select(NavRecord).where(
                NavRecord.document_id == document.id, NavRecord.row_key == row_key
            )
        )
        if prior:
            return prior
    validation = validate_nav(values)
    if values.get("currency") and values["currency"] != product.currency:
        validation.append(
            {
                "rule": "currency_mismatch",
                "message": "币种与产品不符",
                "overridable": False,
            }
        )
    rule = db.get(ValidationRule, manager_id)
    if not validation and rule and rule.max_nav_change is not None:
        previous = db.scalar(
            select(NavRecord)
            .join(EffectiveNav, EffectiveNav.record_id == NavRecord.id)
            .where(
                NavRecord.share_id == share.id,
                NavRecord.valuation_date < values["valuation_date"],
            )
            .order_by(NavRecord.valuation_date.desc())
            .limit(1)
        )
        if previous:
            change = abs(Decimal(values["unit_nav"]) / Decimal(previous.unit_nav) - 1)
            if change > rule.max_nav_change:
                validation.append(
                    {
                        "rule": "nav_change_threshold",
                        "message": "较上一有效估值日的净值变化超出管理员配置阈值",
                        "overridable": True,
                        "change": str(change),
                        "threshold": str(rule.max_nav_change),
                        "base_record_id": previous.id,
                    }
                )
    nav = NavRecord(
        manager_id=manager_id,
        product_id=product.id,
        share_id=share.id,
        valuation_date=values["valuation_date"],
        unit_nav=values["unit_nav"],
        accumulated_nav=values.get("accumulated_nav"),
        net_assets=values.get("net_assets"),
        total_shares=values.get("total_shares"),
        source="manual" if not document or manual else document.source,
        document_id=document.id if document else None,
        actor_id=actor.id if actor else None,
        row_key=row_key,
        received_at=document.received_at if document else now(),
        validation=validation,
        reported_metrics={
            k: values[k]
            for k in ["cash", "position_ratio", "return_rate", "drawdown"]
            if values.get(k) is not None
        },
    )
    db.add(nav)
    db.flush()
    effective = db.get(EffectiveNav, (share.id, nav.valuation_date))
    if validation:
        task(
            db,
            manager_id,
            "validation",
            f"validation:{nav.id}",
            {"record_ids": [nav.id], "errors": validation},
            product.id,
            share.id,
            nav.valuation_date,
        )
    elif effective:
        old = db.get(NavRecord, effective.record_id)
        fields = ["unit_nav", "accumulated_nav", "net_assets", "total_shares"]

        def as_decimal(value):
            return Decimal(value) if value is not None else None

        if (
            any(
                as_decimal(getattr(old, key)) != as_decimal(getattr(nav, key))
                for key in fields
            )
            or old.reported_metrics != nav.reported_metrics
        ):
            candidates = list(
                db.scalars(
                    select(NavRecord.id).where(
                        NavRecord.share_id == share.id,
                        NavRecord.valuation_date == nav.valuation_date,
                    )
                )
            )
            task(
                db,
                manager_id,
                "conflict",
                f"conflict:{share.id}:{nav.valuation_date}",
                {"record_ids": candidates},
                product.id,
                share.id,
                nav.valuation_date,
            )
    else:
        db.add(
            EffectiveNav(
                manager_id=manager_id,
                share_id=share.id,
                valuation_date=nav.valuation_date,
                record_id=nav.id,
            )
        )
    audit(
        db,
        actor,
        manager_id,
        "nav.received",
        nav.id,
        {
            "product_id": product.id,
            "share_id": share.id,
            "valuation_date": nav.valuation_date,
            "validation": validation,
        },
    )
    from .receipts import nav_received
    nav_received(db, nav, actor)
    # Receipt resolves missing only; validation/conflict tasks remain independent.
    missing = db.scalar(
        select(ExceptionTask).where(
            ExceptionTask.manager_id == manager_id,
            ExceptionTask.dedup_key == f"missing:{share.id}:{nav.valuation_date}",
            ExceptionTask.status != "resolved",
        )
    )
    if missing:
        missing.status, missing.resolution, missing.updated_at = (
            "resolved",
            {"via": "data_received", "record_id": nav.id},
            now(),
        )
        missing.revision += 1
        audit(
            db,
            actor,
            manager_id,
            "exception.receipt_resolved",
            missing.id,
            missing.resolution,
        )
    return nav


def organize_nav_document(db, document, record_ids, actor=None):
    """Turn a successfully parsed NAV attachment into organized material metadata."""
    if db.get(DocumentMaterial, document.id):
        return False
    records = list(
        db.scalars(
            select(NavRecord).where(
                NavRecord.manager_id == document.manager_id,
                NavRecord.id.in_(record_ids),
            )
        )
    )
    if not records:
        return False
    product_ids = sorted({record.product_id for record in records})
    valuation_dates = sorted({record.valuation_date for record in records})
    timestamp = now()
    material = DocumentMaterial(
        document_id=document.id,
        manager_id=document.manager_id,
        category="nav_valuation",
        material_type=None,
        title=document.filename,
        business_date=valuation_dates[0] if len(valuation_dates) == 1 else None,
        period_start=valuation_dates[0] if len(valuation_dates) > 1 else None,
        period_end=valuation_dates[-1] if len(valuation_dates) > 1 else None,
        notes="净值解析完成后由系统自动整理；原始文件保持不变。",
        sensitivity="standard",
        status="organized",
        organization_source="system_parse",
        confirmed_by=actor.id if actor else None,
        confirmed_at=timestamp,
        created_at=timestamp,
        updated_at=timestamp,
        revision=1,
    )
    db.add(material)
    db.flush()
    for product_id in product_ids:
        db.add(
            DocumentMaterialProduct(
                manager_id=document.manager_id,
                document_id=document.id,
                product_id=product_id,
                created_at=timestamp,
            )
        )
    audit(
        db,
        actor,
        document.manager_id,
        "document.material_auto_organized",
        document.id,
        {
            "category": "nav_valuation",
            "product_ids": product_ids,
            "business_date": material.business_date,
            "period_start": material.period_start,
            "period_end": material.period_end,
            "record_count": len(records),
            "organization_source": "system_parse",
        },
    )
    return True


def organize_completed_nav_documents(db, limit=500):
    """Backfill NAV files parsed before automatic material organization existed."""
    jobs = list(
        db.scalars(
            select(ParseJob)
            .where(
                ParseJob.status == "completed",
                ~ParseJob.document_id.in_(select(DocumentMaterial.document_id)),
            )
            .order_by(ParseJob.updated_at)
            .limit(limit)
        )
    )
    organized = 0
    for job in jobs:
        document = db.get(Document, job.document_id)
        if document and organize_nav_document(
            db, document, (job.result or {}).get("record_ids", [])
        ):
            organized += 1
    return organized


def process_document(db, settings, document, actor=None):
    job = db.scalar(
        select(ParseJob).where(ParseJob.document_id == document.id).with_for_update()
    )
    if not job:
        job = ParseJob(manager_id=document.manager_id, document_id=document.id)
        db.add(job)
        db.flush()
    job.status, job.updated_at = "processing", now()
    path = settings.storage / document.storage_key
    try:
        metadata = dict(document.metadata_json or {})
        if document.parent_id:
            parent = db.get(Document, document.parent_id)
            if parent:
                metadata = {**(parent.metadata_json or {}), **metadata}
        sender = metadata.get("from", "")
        domain_match = re.search(r"@([A-Za-z0-9.-]+)", sender)
        context = {
            "sender": sender,
            "sender_domain": domain_match.group(1).lower() if domain_match else "",
            "subject": metadata.get("subject", ""),
            "source": document.source,
        }
        result = parse(
            document.filename, path.read_bytes(), settings.max_rows, context=context
        )
    except Exception as exc:
        # Preserve raw evidence and expose only bounded parser diagnostics.
        result = {
            "records": [],
            "errors": [{"reason": str(exc)[:300]}],
            "parser_version": "explicit-header-v2",
        }
    imported = []
    for item in result["records"]:
        query = select(Product).where(Product.manager_id == document.manager_id)
        matched_by_name = False
        if document.product_id:
            query = query.where(Product.id == document.product_id)
        elif item.get("product_code"):
            query = query.where(Product.code == item["product_code"])
        else:
            query = query.where(Product.name == item.get("product_name"))
        matches = list(db.scalars(query))
        if (
            not matches
            and not document.product_id
            and item.get("product_name")
            and result.get("parser_version", "").startswith("custodian:")
        ):
            matches = list(
                db.scalars(
                    select(Product).where(
                        Product.manager_id == document.manager_id,
                        Product.name == item["product_name"],
                    )
                )
            )
            matched_by_name = len(matches) == 1
        if len(matches) != 1:
            result["errors"].append(
                {
                    "row": item["row_key"],
                    "reason": "产品未建档或名称不唯一，需人工确认产品归属",
                    "candidate": item,
                }
            )
            continue
        product = matches[0]
        if (
            item.get("product_code")
            and product.code != item["product_code"]
            and not matched_by_name
        ):
            result["errors"].append(
                {
                    "row": item["row_key"],
                    "reason": "附件产品代码与指定产品不符，禁止强行关联",
                }
            )
            continue
        if not item.get("product_code") and item.get("product_name") != product.name:
            result["errors"].append(
                {
                    "row": item["row_key"],
                    "reason": "附件产品名称与指定产品不符，禁止强行关联",
                }
            )
            continue
        shares = list(
            db.scalars(select(ShareClass).where(ShareClass.product_id == product.id))
        )
        if item.get("share_class"):
            shares = [s for s in shares if s.name == item["share_class"]]
        if len(shares) != 1:
            result["errors"].append(
                {"row": item["row_key"], "reason": "份额无法唯一识别，需先维护产品份额"}
            )
            continue
        # Cumulative-history attachments repeat earlier rows on every delivery.
        # Reuse an identical canonical record instead of creating thousands of
        # duplicate candidates; a changed value still becomes a new conflict.
        def same_number(left, right):
            return (Decimal(left) if left is not None else None) == (
                Decimal(right) if right is not None else None
            )

        identical = next(
            (
                prior
                for prior in db.scalars(
                    select(NavRecord).where(
                        NavRecord.manager_id == document.manager_id,
                        NavRecord.product_id == product.id,
                        NavRecord.share_id == shares[0].id,
                        NavRecord.valuation_date == item["valuation_date"],
                    )
                )
                if all(
                    same_number(getattr(prior, field), item.get(field))
                    for field in [
                        "unit_nav",
                        "accumulated_nav",
                        "net_assets",
                        "total_shares",
                    ]
                )
            ),
            None,
        )
        if identical:
            imported.append(identical.id)
            continue
        nav = add_nav(
            db,
            document.manager_id,
            product,
            shares[0],
            item,
            actor,
            document,
            item["row_key"],
        )
        imported.append(nav.id)
    job.result = {**result, "record_ids": imported}
    job.status, job.updated_at = ("review" if result["errors"] else "completed"), now()
    if job.status == "completed":
        organize_nav_document(db, document, imported, actor)
    if result["errors"]:
        task(
            db,
            document.manager_id,
            "parse",
            f"parse:{document.id}",
            {"document_id": document.id, "job_id": job.id, "errors": result["errors"]},
            document.product_id,
        )
    else:
        old = db.scalar(
            select(ExceptionTask).where(
                ExceptionTask.manager_id == document.manager_id,
                ExceptionTask.dedup_key == f"parse:{document.id}",
            )
        )
        if old and old.status != "resolved":
            old.status, old.resolution, old.updated_at = (
                "resolved",
                {"via": "reparse", "record_ids": imported},
                now(),
            )
            old.revision += 1
            audit(
                db,
                actor,
                document.manager_id,
                "exception.parse_resolved",
                old.id,
                old.resolution,
            )
    audit(
        db,
        actor,
        document.manager_id,
        "document.parsed",
        document.id,
        {"status": job.status, "records": len(imported)},
    )
    return job


def select_effective(db, user, issue, record_id, reversal, reason, revision):
    if issue.status == "resolved":
        raise HTTPException(409, "该待办已经解决")
    record = db.get(NavRecord, record_id)
    if (
        not record
        or record.manager_id != issue.manager_id
        or record.share_id != issue.share_id
        or record.valuation_date != issue.valuation_date
        or record_id not in issue.payload.get("record_ids", [])
    ):
        raise HTTPException(422, "该记录不属于此异常")
    if any(not v.get("overridable") for v in record.validation):
        raise HTTPException(422, "记录有不可豁免的校验错误，需提交正确数据")
    if (reversal or issue.kind == "validation") and not reason.strip():
        raise HTTPException(422, "反账或异常接受必须填写原因")
    db.scalar(
        select(ShareClass).where(ShareClass.id == record.share_id).with_for_update()
    )
    changed = db.execute(
        update(ExceptionTask)
        .where(
            ExceptionTask.id == issue.id,
            ExceptionTask.revision == revision,
            ExceptionTask.status != "resolved",
        )
        .values(
            status="resolved",
            updated_at=now(),
            revision=revision + 1,
            resolution={"record_id": record.id, "reversal": reversal, "reason": reason},
        )
    ).rowcount
    if not changed:
        raise HTTPException(409, "待办已更新，请刷新后处理")
    effective = db.get(EffectiveNav, (record.share_id, record.valuation_date))
    old_id = effective.record_id if effective else None
    if not effective:
        effective = EffectiveNav(
            manager_id=record.manager_id,
            share_id=record.share_id,
            valuation_date=record.valuation_date,
            record_id=record.id,
        )
        db.add(effective)
    else:
        effective.record_id, effective.revision = record.id, effective.revision + 1
    effective.reversal = reversal
    audit(
        db,
        user,
        record.manager_id,
        "nav.reversal" if reversal else "nav.selected",
        record.id,
        {
            "before_record_id": old_id,
            "after_record_id": record.id,
            "reason": reason,
            "valuation_date": record.valuation_date,
            "exception_id": issue.id,
        },
    )


def performance(records):
    # Published accumulated NAV is not assumed to be a reinvested total-return index.
    peak = None
    result = []
    for record in records:
        value = Decimal(record.unit_nav)
        peak = max(peak, value) if peak is not None else value
        result.append(
            {
                "date": record.valuation_date,
                "nav": str(value),
                "nav_change": str(value / Decimal(records[0].unit_nav) - 1),
                "nav_drawdown": str(value / peak - 1),
            }
        )
    return result


def refresh_missing(db, manager_id, on_date=None, at=None):
    from .receipts import check_receipts
    return check_receipts(db, manager_id, on_date=on_date, at=at)
