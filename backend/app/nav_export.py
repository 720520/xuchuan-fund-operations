"""Excel export for the current effective NAV ledger."""

from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SOURCE_LABELS = {
    "manual": "人工录入",
    "upload": "手动上传",
    "email": "托管邮件",
}


def _safe_text(value: object | None) -> str:
    """Keep text-like spreadsheet values from being interpreted as formulas."""
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _excel_datetime(value: object | None) -> datetime | str | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value))
    except ValueError:
        return _safe_text(value)
    if parsed.tzinfo:
        parsed = parsed.astimezone(ZoneInfo("Asia/Shanghai")).replace(tzinfo=None)
    return parsed


def build_nav_export(
    *,
    manager_name: str,
    start_date: date,
    end_date: date,
    records: list[dict],
    generated_at: datetime,
) -> bytes:
    """Build a typed, filterable workbook from already-authorized NAV rows."""
    book = Workbook()
    sheet = book.active
    sheet.title = "产品净值"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D6"

    clay = "A9654C"
    charcoal = "393A35"
    warm = "F5F1EA"
    light = "E7E1D7"
    muted = "6E7068"
    white = "FFFFFF"
    thin = Side(style="thin", color=light)

    sheet.merge_cells("A1:O1")
    title = sheet["A1"]
    title.value = "产品净值导出"
    title.font = Font(name="Microsoft YaHei", size=18, bold=True, color=white)
    title.fill = PatternFill("solid", fgColor=clay)
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34

    metadata = [
        f"管理人：{_safe_text(manager_name)}",
        f"估值日：{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}",
        f"导出时间：{generated_at:%Y-%m-%d %H:%M:%S}",
        f"记录数：{len(records)}",
    ]
    metadata_groups = ((1, 3), (4, 7), (8, 11), (12, 15))
    for (start, end), value in zip(metadata_groups, metadata, strict=True):
        sheet.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)
        cell = sheet.cell(2, start, value)
        cell.font = Font(name="Microsoft YaHei", size=10, color=charcoal)
        cell.fill = PatternFill("solid", fgColor=warm)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 25

    sheet.merge_cells("A3:O3")
    note = sheet["A3"]
    note.value = "只导出当前有效版本；反账记录在“有效状态”列标记。净值数据可通过记录编号和原件编号追溯。"
    note.font = Font(name="Microsoft YaHei", size=9, color=muted)
    note.alignment = Alignment(vertical="center")
    sheet.row_dimensions[3].height = 22

    headers = [
        "估值日",
        "产品代码",
        "产品名称",
        "份额类别",
        "币种",
        "单位净值",
        "累计净值",
        "资产净值",
        "总份额",
        "来源",
        "有效状态",
        "系统收到时间",
        "净值记录编号",
        "原件编号",
        "有效版本号",
    ]
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(5, column, header)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=charcoal)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    sheet.row_dimensions[5].height = 28

    for row_number, record in enumerate(records, 6):
        values = [
            date.fromisoformat(record["valuation_date"]),
            _safe_text(record["product_code"]),
            _safe_text(record["product_name"]),
            _safe_text(record["share_name"]),
            _safe_text(record["currency"]),
            record["unit_nav"],
            record["accumulated_nav"],
            record["net_assets"],
            record["total_shares"],
            SOURCE_LABELS.get(record["source"], _safe_text(record["source"])),
            "反账" if record["reversal"] else "正常",
            _excel_datetime(record["received_at"]),
            _safe_text(record["record_id"]),
            _safe_text(record["document_id"]),
            record["revision"],
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Microsoft YaHei", size=9, color=charcoal)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center")
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="FAF8F4")
        sheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
        for column in (6, 7):
            sheet.cell(row_number, column).number_format = "0.0000000000"
        for column in (8, 9):
            sheet.cell(row_number, column).number_format = "#,##0.0000"
        sheet.cell(row_number, 12).number_format = "yyyy-mm-dd hh:mm:ss"

    last_row = max(5, 5 + len(records))
    sheet.auto_filter.ref = f"A5:O{last_row}"
    widths = [13, 18, 38, 14, 9, 16, 16, 20, 20, 14, 14, 21, 38, 38, 15]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:5"
    book.properties.title = "产品净值导出"
    book.properties.subject = f"{start_date:%Y-%m-%d} 至 {end_date:%Y-%m-%d}"
    book.properties.creator = "序川基金运营工作台"

    output = BytesIO()
    book.save(output)
    return output.getvalue()
