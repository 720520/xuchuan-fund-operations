from io import BytesIO

from app.models import AuditEvent, ProductGrant
from conftest import login, nav_data, product
from openpyxl import load_workbook
from sqlalchemy import select


def workbook_from(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content), data_only=False)


def test_export_effective_nav_by_date_and_product_with_audit(env):
    app, client, ids = env
    login(client)
    selected = product(client, ids["a"], "F001")
    other = product(client, ids["a"], "F002")
    client.post(
        f"/api/managers/{ids['a']}/nav",
        json={
            **nav_data(selected, "1.0512", "2026-09-11"),
            "accumulated_nav": "1.2034",
            "total_shares": "1200000.25",
        },
    )
    client.post(
        f"/api/managers/{ids['a']}/nav",
        json=nav_data(selected, "1.0600", "2026-09-12"),
    )
    client.post(
        f"/api/managers/{ids['a']}/nav",
        json=nav_data(other, "2.1000", "2026-09-11"),
    )

    response = client.get(
        f"/api/managers/{ids['a']}/nav/export",
        params={
            "start_date": "2026-09-11",
            "end_date": "2026-09-11",
            "product_id": selected["id"],
        },
    )
    book = workbook_from(response)
    sheet = book["产品净值"]
    assert sheet.freeze_panes == "D6"
    assert sheet.auto_filter.ref == "A5:O6"
    assert sheet["A2"].value == "管理人：测试牌照a"
    assert sheet["D2"].value == "估值日：2026-09-11 至 2026-09-11"
    assert sheet["H2"].value.startswith("导出时间：")
    assert sheet["L2"].value == "记录数：1"
    assert [cell.value for cell in sheet[5]][:7] == [
        "估值日",
        "产品代码",
        "产品名称",
        "份额类别",
        "币种",
        "单位净值",
        "累计净值",
    ]
    assert sheet["A6"].value.strftime("%Y-%m-%d") == "2026-09-11"
    assert sheet["B6"].value == "F001"
    assert sheet["F6"].value == 1.0512
    assert sheet["G6"].value == 1.2034
    assert sheet["I6"].value == 1200000.25
    assert sheet["J6"].value == "人工录入"
    assert sheet["K6"].value == "正常"
    assert sheet.max_row == 6

    with app.state.factory() as session:
        event = session.scalar(
            select(AuditEvent).where(AuditEvent.action == "nav.exported")
        )
        assert event.details == {
            "start_date": "2026-09-11",
            "end_date": "2026-09-11",
            "product_ids": [selected["id"]],
            "scope": "selected",
            "row_count": 1,
        }


def test_export_respects_product_and_download_permissions(env):
    app, client, ids = env
    login(client)
    allowed = product(client, ids["a"], "ALLOW")
    denied = product(client, ids["a"], "DENY")
    for item in (allowed, denied):
        client.post(
            f"/api/managers/{ids['a']}/nav",
            json=nav_data(item, day="2026-09-11"),
        )
    with app.state.factory.begin() as session:
        session.add(ProductGrant(user_id=ids["fund"], product_id=allowed["id"]))

    login(client, "fund")
    response = client.get(
        f"/api/managers/{ids['a']}/nav/export",
        params={"start_date": "2026-09-11", "end_date": "2026-09-11"},
    )
    sheet = workbook_from(response)["产品净值"]
    assert sheet.max_row == 6
    assert sheet["B6"].value == "ALLOW"
    assert (
        client.get(
            f"/api/managers/{ids['a']}/nav/export",
            params={
                "start_date": "2026-09-11",
                "end_date": "2026-09-11",
                "product_id": denied["id"],
            },
        ).status_code
        == 403
    )

    login(client, "colleague")
    assert (
        client.get(
            f"/api/managers/{ids['a']}/nav/export",
            params={"start_date": "2026-09-11", "end_date": "2026-09-11"},
        ).status_code
        == 403
    )
    login(client)
    assert (
        client.get(
            f"/api/managers/{ids['a']}/nav/export",
            params={"start_date": "2026-09-12", "end_date": "2026-09-11"},
        ).status_code
        == 422
    )


def test_export_multiple_selected_products_in_one_sheet(env):
    _, client, ids = env
    login(client)
    products = [product(client, ids["a"], code) for code in ("P001", "P002", "P003")]
    for index, item in enumerate(products, 1):
        client.post(
            f"/api/managers/{ids['a']}/nav",
            json=nav_data(item, value=f"1.{index:02d}", day="2026-09-11"),
        )

    response = client.get(
        f"/api/managers/{ids['a']}/nav/export",
        params=[
            ("start_date", "2026-09-11"),
            ("end_date", "2026-09-11"),
            ("product_id", products[0]["id"]),
            ("product_id", products[2]["id"]),
        ],
    )
    sheet = workbook_from(response)["产品净值"]
    assert sheet["L2"].value == "记录数：2"
    assert sheet.max_row == 7
    assert {sheet["B6"].value, sheet["B7"].value} == {"P001", "P003"}


def test_export_escapes_spreadsheet_formula_like_labels(env):
    _, client, ids = env
    login(client)
    created = client.post(
        f"/api/managers/{ids['a']}/products",
        json={"code": "=F001", "name": "+产品", "shares": ["@A"]},
    )
    assert created.status_code == 201, created.text
    exported_product = next(
        item
        for item in client.get(f"/api/managers/{ids['a']}/products").json()
        if item["id"] == created.json()["id"]
    )
    client.post(
        f"/api/managers/{ids['a']}/nav",
        json=nav_data(exported_product, day="2026-09-11"),
    )
    sheet = workbook_from(
        client.get(
            f"/api/managers/{ids['a']}/nav/export",
            params={"start_date": "2026-09-11", "end_date": "2026-09-11"},
        )
    )["产品净值"]
    assert sheet["B6"].value == "'=F001"
    assert sheet["C6"].value == "'+产品"
    assert sheet["D6"].value == "'@A"
    assert sheet["B6"].data_type == "s"
